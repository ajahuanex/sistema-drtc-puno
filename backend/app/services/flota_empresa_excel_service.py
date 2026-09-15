"""
Servicio de importación masiva Excel para el módulo Flota Empresa.
Procesa el archivo DB_VEHICULOS con normalización de TUC, rutas y observaciones.
Sincroniza automáticamente la Razón Social con 'empresas', las primigenias con 'resoluciones_primigenias'
y crea/actualiza automáticamente las resoluciones hijas en 'resoluciones_hijas'.
"""
import re
import uuid
import logging
from io import BytesIO
from typing import List, Dict, Any, Optional
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.models.flota_empresa import (
    VehiculoEmpresaCreate,
    EntradaObservacion
)

logger = logging.getLogger(__name__)

# -----------------------------------------------------------------------
# Mapeo de columnas del Excel de origen (DB_VEHICULOS)
# -----------------------------------------------------------------------
COLUMNAS_MAP = {
    "A": "ruc",
    "B": "nro_resolucion_primigenia",
    "C": "nro_resolucion_hija",
    # D: porcentaje – ignorado
    "E": "placa",
    "F": "ruta",
    "G": "tuc",
    "H": "estado",
    "I": "observaciones",
    "J": "fecha_cronologica",
    "K": "razon_social",
    "L": "fecha_resolucion_hija",
    "M": "id_origen",
    "N": "notificado",
    "O": "estado_primigenia",
}

ESTADOS_VALIDOS = {"HABILITADO", "INHABILITADO", "OBSERVADO", "CANCELADO", "SUSPENDIDO"}
TIPOS_HIJA_VALIDOS = {"I", "S", "M", "O", "C", "R"}
PLACA_REGEX = re.compile(r"^[A-Z0-9]{1,3}-[A-Z0-9]{3,4}$", re.IGNORECASE)


def _clean_str(val) -> Optional[str]:
    """Limpiar valor de celda a string o None."""
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return None
    s = str(val).strip()
    return s if s and s.upper() not in ("NAN", "NONE", "-", "") else None


def _normalizar_expediente(val) -> Optional[str]:
    """
    Normalizar número de expediente a formato 'E-0123-2026':
    - '0123-2026-E' → 'E-0123-2026'
    - '0123-2026' → 'E-0123-2026'
    - 'EXP-0123-2026' → 'E-0123-2026'
    - 'E-0123-2026' → 'E-0123-2026'
    """
    s = _clean_str(val)
    if not s:
        return None
    s = s.upper().replace(" ", "")
    if s.startswith("EXP-"):
        s = s[4:]
    elif s.startswith("EXP"):
        s = s[3:]
    if s.endswith("-E"):
        s = s[:-2]
    elif s.endswith("E") and len(s) > 1 and s[-2].isdigit():
        s = s[:-1]
    
    if s.startswith("E-"):
        return s
    if s.startswith("E") and len(s) > 1 and s[1].isdigit():
        s = f"E-{s[1:]}"
    return f"E-{s}"


def _normalizar_primigenia(val) -> Optional[str]:
    """
    Normalizar número de resolución primigenia para que siempre comience con 'R-':
    - '0123-2026' → 'R-0123-2026'
    - 'R-0123-2026' → 'R-0123-2026'
    - '0128-2024' → 'R-0128-2024'
    """
    s = _clean_str(val)
    if not s:
        return None
    s = s.upper().replace(" ", "")
    if s.startswith("R-"):
        return s
    return f"R-{s}"


def _normalizar_tuc(val, placa: Optional[str] = None) -> Optional[str]:
    """
    Normalizar número TUC a 8 caracteres 'T-012345' (6 dígitos formateados).
    Si no tiene TUC o es inválido/vacio y se provee placa válida (ej: 'A2B-123'),
    retornar 'T-A2B-123'.
    """
    s = _clean_str(val)
    if not s or s in ("-", "–", "—"):
        if placa and placa != "-":
            return f"T-{placa.upper()}"
        return None

    s_upper = s.upper()
    if s_upper.startswith("T-"):
        sin_prefix = s_upper[2:]
        if sin_prefix.isdigit():
            return f"T-{sin_prefix.zfill(6)}"
        return s_upper

    digits = re.sub(r"[^\d]", "", s)
    if digits:
        return f"T-{digits.zfill(6)}"

    if placa and placa != "-":
        return f"T-{placa.upper()}"
    return s_upper


def _normalizar_rutas(val) -> List[str]:
    """
    Normalizar campo de rutas:
    - '10203' → ['01', '02', '03']
    - '01 02 03' → ['01', '02', '03']
    - '01, 02, 03' → ['01', '02', '03']
    - '01-02-03' → ['01', '02', '03']
    """
    s = _clean_str(val)
    if not s:
        return []
    
    # 1. Si contiene delimitadores (coma, espacio, guion, slash, pipe)
    if re.search(r"[,\s\-/|]", s):
        partes = re.split(r"[,\s\-/|]+", s)
        result = []
        for p in partes:
            p = p.strip()
            if p and p.upper() not in ("NAN", "NONE", "-", ""):
                if p.isdigit():
                    p = p.zfill(2)
                result.append(p)
        return list(dict.fromkeys(result))
    
    # 2. Si NO contiene delimitadores y son solo dígitos (ej: "10203", "010203", "01", "1")
    if s.isdigit():
        if len(s) == 1:
            return [s.zfill(2)]
        if len(s) == 2:
            return [s]
        tokens = []
        i = 0
        if len(s) % 2 != 0:
            tokens.append(s[0].zfill(2))
            i = 1
        while i < len(s):
            pair = s[i:i+2]
            tokens.append(pair.zfill(2))
            i += 2
        return list(dict.fromkeys(tokens))
    
    return [s]


def _normalizar_placa(val) -> tuple[str, bool]:
    """
    Normalizar placa. Retorna (placa, es_cronologico).
    Si vacío/guion → ("-", True).
    """
    s = _clean_str(val)
    if not s or s in ("-", "–", "—"):
        return "-", True
    placa = s.upper().replace(" ", "")
    return placa, False


def _normalizar_estado(val) -> Optional[str]:
    """Normalizar estado vehicular."""
    s = _clean_str(val)
    if not s:
        return None
    s_upper = s.upper()
    mapeo = {
        "HABILITADO": "HABILITADO",
        "INHABILITADO": "INHABILITADO",
        "INHABIILITADO": "INHABILITADO",
        "OBSERVADO": "OBSERVADO",
        "CANCELADO": "CANCELADO",
        "SUSPENDIDO": "SUSPENDIDO",
        "SUSPENDDO": "SUSPENDIDO",
        "-": None,
        "–": None,
    }
    return mapeo.get(s_upper, s_upper if s_upper in ESTADOS_VALIDOS else None)


def _extraer_tipo_hija(nro_hija: Optional[str]) -> Optional[str]:
    """Extraer tipo de resolución hija del número. 0133-2024-S → 'S'."""
    if not nro_hija:
        return None
    m = re.search(r"-([ISMOCR])(?:\s|$)", nro_hija.upper())
    if m:
        return m.group(1)
    partes = nro_hija.upper().split("-")
    if partes and len(partes[-1]) == 1 and partes[-1] in TIPOS_HIJA_VALIDOS:
        return partes[-1]
    return None


def _parse_observaciones(val) -> List[EntradaObservacion]:
    """Parsear observaciones separadas por | en array de EntradaObservacion."""
    s = _clean_str(val)
    if not s:
        return []
    partes = [p.strip() for p in s.split("|") if p.strip()]
    return [EntradaObservacion(texto=p, fuente="importacion") for p in partes]


def _parse_fecha(val) -> Optional[datetime]:
    """Parsear fechas desde diferentes formatos."""
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return None
    if isinstance(val, datetime):
        return val
    s = _clean_str(val)
    if not s:
        return None
    formatos = ["%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d", "%m/%d/%Y"]
    for fmt in formatos:
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


class FlotaEmpresaExcelService:
    def __init__(self, db):
        self.db = db
        self.collection = db["flota_empresa"]
        self.empresas_coll = db["empresas"]
        self.primigenias_coll = db["resoluciones_primigenias"]
        self.hijas_coll = db["resoluciones_hijas"]

    # ------------------------------------------------------------------
    # HELPERS DE CRUCE E INTEGRACIÓN ENTRE MÓDULOS
    # ------------------------------------------------------------------

    async def _obtener_razon_social(self, ruc: Optional[str], razon_excel: Optional[str], cache: dict) -> Optional[str]:
        """Obtener la Razón Social oficial desde la colección 'empresas' usando RUC."""
        if not ruc:
            return razon_excel
        if ruc in cache:
            return cache[ruc] or razon_excel
        
        emp_doc = await self.empresas_coll.find_one({"ruc": ruc})
        if not emp_doc:
            emp_doc = await self.empresas_coll.find_one({"ruc": {"$regex": f"^{re.escape(ruc)}$"}})
        
        razon_oficial = None
        if emp_doc:
            rs = emp_doc.get("razonSocial")
            if isinstance(rs, dict):
                razon_oficial = rs.get("principal") or rs.get("sunat") or rs.get("minimo")
            elif isinstance(rs, str) and rs.strip():
                razon_oficial = rs.strip()
                
            if not razon_oficial:
                datos_sunat = emp_doc.get("datosSunat")
                if isinstance(datos_sunat, dict):
                    razon_oficial = datos_sunat.get("ddp_nombre") or datos_sunat.get("razonSocial")
                    
            if not razon_oficial:
                razon_oficial = emp_doc.get("razon_social") or emp_doc.get("nombre_comercial")
        
        cache[ruc] = razon_oficial
        return razon_oficial or razon_excel

    async def _obtener_primigenia(self, nro_prim: Optional[str], cache: dict) -> Optional[dict]:
        """
        Obtener datos de la Resolución Primigenia desde 'resoluciones_primigenias'.
        Si el número ingresado en la columna B correspondiera en realidad a una Resolución Hija (ej: '0128-2024'),
        se busca en 'resoluciones_hijas' para resolver automáticamente la primigenia matriz padre.
        """
        if not nro_prim:
            return None
        nro_clean = nro_prim.strip()
        nro_r = _normalizar_primigenia(nro_clean) or nro_clean
        
        if nro_r in cache:
            return cache[nro_r]
        if nro_clean in cache:
            return cache[nro_clean]
        
        # 1. Buscar primero en 'resoluciones_primigenias'
        regex_pat = f"^R?-?{re.escape(nro_clean.replace('R-', ''))}$"
        doc = await self.primigenias_coll.find_one({
            "$or": [
                {"nro_resolucion": nro_r},
                {"nro_resolucion": nro_clean},
                {"nro_resolucion": {"$regex": regex_pat, "$options": "i"}}
            ]
        })
        
        # 2. Si NO se encuentra como primigenia, verificar si es una Resolución Hija (ej: 0128-2024 es Hija Sustitución)
        if not doc:
            hija_doc = await self.hijas_coll.find_one({
                "$or": [
                    {"nro_resolucion": nro_clean},
                    {"nro_resolucion": nro_r},
                    {"nro_resolucion": {"$regex": f"^{re.escape(nro_clean)}$", "$options": "i"}}
                ]
            })
            if hija_doc and hija_doc.get("nro_resolucion_primigenia"):
                padre_nro = _normalizar_primigenia(hija_doc.get("nro_resolucion_primigenia"))
                if padre_nro:
                    doc = await self.primigenias_coll.find_one({
                        "$or": [
                            {"nro_resolucion": padre_nro},
                            {"nro_resolucion": {"$regex": f"^R?-?{re.escape(padre_nro.replace('R-', ''))}$", "$options": "i"}}
                        ]
                    })
                    if doc:
                        doc["_resolucion_hija_detectada"] = nro_clean
                        doc["_nro_primigenia_resuelta"] = padre_nro
        
        cache[nro_r] = doc
        cache[nro_clean] = doc
        return doc

    async def _sincronizar_resolucion_hija(
        self,
        nro_hija: Optional[str],
        nro_prim: str,
        tipo_hija_code: Optional[str],
        ruc: str,
        razon_social: Optional[str],
        placa: str,
        fecha_hija: Optional[datetime],
        fecha_crono: Optional[datetime],
        prim_doc: Optional[dict],
        hija_cache: dict
    ):
        """
        Validar y sincronizar la resolución hija en el módulo 'resoluciones_hijas'.
        Si la resolución primigenia no tuviera la hija registrada, se importa/crea
        o actualiza automáticamente en ese módulo.
        """
        if not nro_hija:
            return
        
        nro_hija_clean = nro_hija.strip()
        now = datetime.utcnow()
        fecha_efecto = fecha_hija or fecha_crono or now
        
        mapeo_tipo = {
            "I": "INCREMENTO_FLOTA",
            "S": "SUSTITUCION_VEHICULAR",
            "M": "MODIFICACION_RUTA",
            "O": "OTROS",
            "C": "CANCELACION_PARCIAL",
            "R": "RENOVACION",
        }
        tipo_acto = mapeo_tipo.get((tipo_hija_code or "").upper(), "INCREMENTO_FLOTA" if (placa and placa != "-") else "OTROS")
        
        hija_doc = None
        if nro_hija_clean in hija_cache:
            hija_doc = hija_cache[nro_hija_clean]
        else:
            hija_doc = await self.hijas_coll.find_one({"nro_resolucion": nro_hija_clean})
            if not hija_doc:
                hija_doc = await self.hijas_coll.find_one({
                    "nro_resolucion": {"$regex": f"^{re.escape(nro_hija_clean)}$", "$options": "i"}
                })
        
        if not hija_doc:
            # CREAR AUTOMÁTICAMENTE la Resolución Hija en resoluciones_hijas
            hija_id = str(uuid.uuid4())
            nueva_hija = {
                "id": hija_id,
                "nro_resolucion": nro_hija_clean,
                "nro_resolucion_primigenia": nro_prim,
                "resolucion_primigenia_id": prim_doc.get("id") or str(prim_doc.get("_id")) if prim_doc else None,
                "ruc_empresa": ruc,
                "razon_social": razon_social,
                "tipo_acto": tipo_acto,
                "fecha_resolucion": fecha_efecto,
                "fecha_inicio_efectos": fecha_efecto,
                "vehiculos_ingresantes": [placa] if (placa and placa != "-") else [],
                "vehiculos_salientes": [],
                "rutas_modificadas_ids": [],
                "observaciones": f"Importado automáticamente desde Carga Masiva Flota Vehicular (RUC: {ruc})",
                "esta_activo": True,
                "fecha_registro": now,
                "fecha_actualizacion": now,
            }
            await self.hijas_coll.insert_one(nueva_hija)
            hija_cache[nro_hija_clean] = nueva_hija
            
            # Si existe la resolución primigenia, vincular al historial_modificaciones
            if prim_doc and "_id" in prim_doc:
                mod_entry = {
                    "resolucion_hija_id": hija_id,
                    "nro_resolucion_hija": nro_hija_clean,
                    "tipo_modificacion": tipo_acto,
                    "fecha_acto": fecha_efecto,
                    "observacion": f"Importado automáticamente de Flota (Placa {placa})"
                }
                await self.primigenias_coll.update_one(
                    {"_id": prim_doc["_id"]},
                    {
                        "$push": {"historial_modificaciones": mod_entry},
                        "$set": {"fecha_actualizacion": now}
                    }
                )
        else:
            # ACTUALIZAR Resolución Hija existente
            hija_cache[nro_hija_clean] = hija_doc
            updates = {"fecha_actualizacion": now}
            if ruc and not hija_doc.get("ruc_empresa"):
                updates["ruc_empresa"] = ruc
            if razon_social and not hija_doc.get("razon_social"):
                updates["razon_social"] = razon_social
            if prim_doc and not hija_doc.get("resolucion_primigenia_id"):
                updates["resolucion_primigenia_id"] = prim_doc.get("id") or str(prim_doc.get("_id"))
                
            push_updates = {}
            if placa and placa != "-":
                veh_ingresantes = hija_doc.get("vehiculos_ingresantes", [])
                if placa not in veh_ingresantes:
                    push_updates["vehiculos_ingresantes"] = placa
                    
            op = {"$set": updates}
            if push_updates:
                op["$addToSet"] = push_updates
                
            await self.hijas_coll.update_one({"_id": hija_doc["_id"]}, op)

    # ------------------------------------------------------------------
    # VALIDACIÓN Y PREVIEW
    # ------------------------------------------------------------------

    def _procesar_fila(self, idx: int, row: pd.Series, columnas: list) -> dict:
        """Procesar una fila del DataFrame y retornar dict con datos y errores."""
        errores = []

        def get_col(pos_letter: str, alt_names: list = None):
            # 1. Intentar por nombres de columna alternativos (flexibles)
            if alt_names:
                for name in alt_names:
                    # Coincidencia exacta o flexible
                    for col_name in row.index:
                        col_str = str(col_name).strip().upper()
                        target_str = str(name).strip().upper()
                        if col_str == target_str or col_str.replace(" ", "_") == target_str.replace(" ", "_"):
                            val = row[col_name]
                            if val is not None and not (isinstance(val, float) and pd.isna(val)):
                                return val

            # 2. Intentar por posición de letra (A=0, B=1, etc.)
            col_idx = ord(pos_letter.upper()) - ord("A")
            if col_idx < len(columnas):
                val = row.iloc[col_idx] if hasattr(row, 'iloc') else None
                if val is not None and not (isinstance(val, float) and pd.isna(val)):
                    return val
            
            # 3. Respaldo por índice numérico si el objeto row admite indexación entera
            try:
                val = row[col_idx]
                if val is not None and not (isinstance(val, float) and pd.isna(val)):
                    return val
            except Exception:
                pass
            return None

        # ---- Campos obligatorios ----
        ruc_raw = get_col("A", ["RUC", "ruc"])
        ruc = _clean_str(ruc_raw)
        if not ruc or len(ruc) < 8:
            errores.append(f"RUC inválido: '{ruc_raw}'")

        nro_prim_raw = get_col("B", ["RDR_PRIMIGENIA", "RDR PRIMIGENIA", "RDR_PRIMIGENIA ", "nro_resolucion_primigenia", "RESOLUCION_PRIMIGENIA", "RESOLUCION PRIMIGENIA", "PRIMIGENIA"])
        nro_prim = _normalizar_primigenia(nro_prim_raw)
        if not nro_prim:
            errores.append("Sin resolución primigenia (columna B)")

        # ---- Campos opcionales ----
        nro_hija = _clean_str(get_col("C", ["RDR", "nro_resolucion_hija"]))
        tipo_hija = _extraer_tipo_hija(nro_hija)

        placa, es_cronologico = _normalizar_placa(get_col("E", ["PLACA", "placa"]))
        rutas = _normalizar_rutas(get_col("F", ["RUTA", "ruta"]))
        tuc = _normalizar_tuc(get_col("G", ["TUC", "tuc"]), placa=placa)
        estado = _normalizar_estado(get_col("H", ["ESTADO", "estado"]))
        obs = _parse_observaciones(get_col("I", ["OBSERVACIONES", "observaciones"]))
        fecha_crono = _parse_fecha(get_col("J", ["FECHA", "fecha_cronologica"]))
        razon_social = _clean_str(get_col("K", ["RAZON SOCIAL", "razon_social"]))
        fecha_hija = _parse_fecha(get_col("L", ["FECHA HIJA", "fecha_resolucion_hija"]))
        id_origen = _clean_str(get_col("M", ["ID", "id_origen"]))
        notificado = _clean_str(get_col("N", ["NOTIFICADO", "notificado"]))
        estado_prim = _clean_str(get_col("O", ["ESTADO PRIMIGENIA", "estado_primigenia"]))

        num_expediente = _normalizar_expediente(get_col("P", ["NUM_EXPEDIENTE", "EXPEDIENTE", "num_expediente"]))
        fecha_expediente = _parse_fecha(get_col("Q", ["FECHA_EXPEDIENTE", "FECHA EXPEDIENTE", "fecha_expediente"]))
        link_tuc = _clean_str(get_col("R", ["LINK_TUC", "LINK TUC", "link_tuc"]))
        link_notificacion = _clean_str(get_col("S", ["LINK_NOTIFICACION", "LINK NOTIFICACION", "link_notificacion"]))
        detalles = _clean_str(get_col("T", ["DETALLES", "detalles"]))

        return {
            "fila": idx + 2,
            "ruc": ruc,
            "nro_resolucion_primigenia": nro_prim,
            "nro_resolucion_hija": nro_hija,
            "tipo_resolucion_hija": tipo_hija,
            "placa": placa,
            "es_cronologico": es_cronologico,
            "rutas": rutas,
            "numero_tuc": tuc,
            "estado": estado,
            "observaciones_historial": obs,
            "fecha_cronologica": fecha_crono,
            "razon_social": razon_social,
            "fecha_resolucion_hija": fecha_hija,
            "id_origen": id_origen,
            "notificado": notificado,
            "estado_primigenia": estado_prim,
            "num_expediente": num_expediente,
            "fecha_expediente": fecha_expediente,
            "link_tuc": link_tuc,
            "link_notificacion": link_notificacion,
            "detalles": detalles,
            "es_valido": len(errores) == 0,
            "errores": errores,
        }

    async def procesar_preview(self, buffer: BytesIO, n_filas: int = 15) -> List[dict]:
        """Procesar solo las primeras N filas para preview con cruce de módulos."""
        try:
            df = pd.read_excel(buffer, header=0, dtype=str)
            df = df.fillna("")
        except Exception:
            buffer.seek(0)
            df = pd.read_csv(buffer, header=0, dtype=str, encoding="utf-8-sig")
            df = df.fillna("")

        columnas = list(df.columns)
        resultados = []
        emp_cache = {}
        prim_cache = {}

        for idx, row in df.head(n_filas).iterrows():
            try:
                r = self._procesar_fila(idx, row, columnas)
                # Cruce con módulo Empresas
                r["razon_social"] = await self._obtener_razon_social(r.get("ruc"), r.get("razon_social"), emp_cache)
                # Cruce con módulo Resoluciones Primigenias
                prim_doc = await self._obtener_primigenia(r.get("nro_resolucion_primigenia"), prim_cache)
                if prim_doc and prim_doc.get("estado"):
                    r["estado_primigenia"] = prim_doc.get("estado")
                resultados.append(r)
            except Exception as e:
                resultados.append({"fila": idx + 2, "es_valido": False, "errores": [str(e)]})
        return resultados

    # ------------------------------------------------------------------
    # CARGA MASIVA REAL (CON CRUCE E INTEGRACIÓN AUTOMÁTICA)
    # ------------------------------------------------------------------

    async def procesar_carga_masiva(self, buffer: BytesIO, modo: str = "upsert") -> Dict[str, Any]:
        """
        Importar el archivo a la colección flota_empresa con:
        1. Razón social obtenida de 'empresas'.
        2. Resolución primigenia enlazada con 'resoluciones_primigenias'.
        3. Registro/actualización automática de resoluciones hijas en 'resoluciones_hijas'.
        """
        try:
            df = pd.read_excel(buffer, header=0, dtype=str)
            df = df.fillna("")
        except Exception:
            buffer.seek(0)
            df = pd.read_csv(buffer, header=0, dtype=str, encoding="utf-8-sig")
            df = df.fillna("")

        columnas = list(df.columns)
        total = len(df)
        creados = 0
        actualizados = 0
        omitidos = 0
        errores_list = []

        emp_cache = {}
        prim_cache = {}
        hija_cache = {}

        for idx, row in df.iterrows():
            try:
                datos = self._procesar_fila(idx, row, columnas)
                if not datos["es_valido"] or not datos["ruc"] or not datos["nro_resolucion_primigenia"]:
                    omitidos += 1
                    errores_list.append({
                        "fila": datos["fila"],
                        "error": "; ".join(datos.get("errores", ["Datos inválidos"]))
                    })
                    continue

                # 1. Traer Razón Social del módulo Empresas
                datos["razon_social"] = await self._obtener_razon_social(datos["ruc"], datos["razon_social"], emp_cache)

                # 2. Traer datos de la Primigenia del módulo Resoluciones Primigenias (o resolver si B era resolución hija)
                prim_doc = await self._obtener_primigenia(datos["nro_resolucion_primigenia"], prim_cache)
                if prim_doc:
                    if prim_doc.get("estado"):
                        datos["estado_primigenia"] = prim_doc.get("estado")
                    if prim_doc.get("fecha_vigencia_hasta"):
                        datos["fecha_vigencia_hasta"] = prim_doc.get("fecha_vigencia_hasta")
                    if prim_doc.get("_nro_primigenia_resuelta"):
                        # Si la resolución en columna B era una resolución hija (ej: 0128-2024), asignar su verdadera primigenia matriz (ej: R-0576-2022)
                        # y preservar el número de resolución hija si no venía en columna C
                        if not datos.get("nro_resolucion_hija") and prim_doc.get("_resolucion_hija_detectada"):
                            datos["nro_resolucion_hija"] = prim_doc.get("_resolucion_hija_detectada")
                        datos["nro_resolucion_primigenia"] = prim_doc.get("_nro_primigenia_resuelta")

                # 3. Sincronizar/Auto-crear Resolución Hija en el módulo Resoluciones Hijas
                if datos.get("nro_resolucion_hija"):
                    await self._sincronizar_resolucion_hija(
                        nro_hija=datos["nro_resolucion_hija"],
                        nro_prim=datos["nro_resolucion_primigenia"],
                        tipo_hija_code=datos.get("tipo_resolucion_hija"),
                        ruc=datos["ruc"],
                        razon_social=datos["razon_social"],
                        placa=datos["placa"],
                        fecha_hija=datos.get("fecha_resolucion_hija"),
                        fecha_crono=datos.get("fecha_cronologica"),
                        prim_doc=prim_doc,
                        hija_cache=hija_cache
                    )

                # 4. Construir documento para flota_empresa
                now = datetime.utcnow()
                doc_data = {
                    "ruc": datos["ruc"],
                    "razon_social": datos["razon_social"],
                    "nro_resolucion_primigenia": datos["nro_resolucion_primigenia"],
                    "nro_resolucion_hija": datos["nro_resolucion_hija"],
                    "tipo_resolucion_hija": datos["tipo_resolucion_hija"],
                    "placa": datos["placa"],
                    "es_cronologico": datos["es_cronologico"],
                    "rutas": datos["rutas"],
                    "numero_tuc": datos["numero_tuc"],
                    "estado": datos["estado"],
                    "observaciones_historial": [
                        o.model_dump() for o in datos["observaciones_historial"]
                    ],
                    "fecha_cronologica": datos["fecha_cronologica"],
                    "fecha_resolucion_hija": datos["fecha_resolucion_hija"],
                    "id_origen": datos["id_origen"],
                    "notificado": datos["notificado"],
                    "estado_primigenia": datos["estado_primigenia"],
                    "num_expediente": datos.get("num_expediente"),
                    "fecha_expediente": datos.get("fecha_expediente"),
                    "link_tuc": datos.get("link_tuc"),
                    "link_notificacion": datos.get("link_notificacion"),
                    "detalles": datos.get("detalles"),
                    "esta_activo": True,
                }

                if modo == "upsert":
                    filtro_upsert = {
                        "ruc": datos["ruc"],
                        "nro_resolucion_primigenia": datos["nro_resolucion_primigenia"],
                        "nro_resolucion_hija": datos["nro_resolucion_hija"],
                        "placa": datos["placa"],
                    }
                    existente = await self.collection.find_one(filtro_upsert)
                    if existente:
                        doc_data["fecha_actualizacion"] = now
                        await self.collection.update_one(
                            {"_id": existente["_id"]},
                            {"$set": doc_data}
                        )
                        actualizados += 1
                    else:
                        doc_data["fecha_registro"] = now
                        doc_data["fecha_actualizacion"] = now
                        await self.collection.insert_one(doc_data)
                        creados += 1
                else:
                    doc_data["fecha_registro"] = now
                    doc_data["fecha_actualizacion"] = now
                    await self.collection.insert_one(doc_data)
                    creados += 1

            except Exception as e:
                logger.error(f"Error procesando fila {idx+2}: {e}", exc_info=True)
                omitidos += 1
                errores_list.append({"fila": idx + 2, "error": str(e)})

        return {
            "total_filas": total,
            "creados": creados,
            "actualizados": actualizados,
            "omitidos": omitidos,
            "errores": errores_list[:50],
            "modo": modo,
        }

    # ------------------------------------------------------------------
    # PLANTILLA EXCEL
    # ------------------------------------------------------------------

    def generar_plantilla_excel(self) -> BytesIO:
        """Generar plantilla Excel para carga masiva de flota empresa."""
        wb = Workbook()
        ws = wb.active
        ws.title = "FLOTA_EMPRESA"

        header_fill = PatternFill("solid", fgColor="1E3A5F")
        header_font = Font(color="FFFFFF", bold=True, size=10)
        border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin")
        )
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)

        headers = [
            ("RUC", "A", True, "RUC de la empresa (11 dígitos)"),
            ("RDR_PRIMIGENIA", "B", True, "N° Resolución Primigenia (ej: R-0128-2024)"),
            ("RDR", "C", False, "N° Resolución Hija: ej: 0133-2024-S"),
            ("PORCENTAJE", "D", False, "Porcentaje (opcional)"),
            ("PLACA", "E", True, "Placa: A2B-123. Dejar guion (-) si no aplica"),
            ("RUTA", "F", False, "Códigos de ruta separados por coma: 01,02,03"),
            ("TUC", "G", False, "N° TUC oficial (ej: T-010145 o T-A2B-123)"),
            ("ESTADO", "H", False, "HABILITADO|INHABILITADO|OBSERVADO|CANCELADO|SUSPENDIDO"),
            ("OBSERVACIONES", "I", False, "Observaciones separadas por | para historial"),
            ("FECHA", "J", False, "Fecha cronológica (DD/MM/YYYY)"),
            ("RAZON SOCIAL", "K", False, "Razón social de la empresa"),
            ("FECHA HIJA", "L", False, "Fecha de emisión de resolución hija (DD/MM/YYYY)"),
            ("ID", "M", False, "ID externo de origen"),
            ("NOTIFICADO", "N", False, "Indicador de notificación"),
            ("ESTADO PRIMIGENIA", "O", False, "Estado de la resolución primigenia: ACTIVA/INACTIVA"),
            ("NUM_EXPEDIENTE", "P", False, "Número de Expediente Administrativo"),
            ("FECHA_EXPEDIENTE", "Q", False, "Fecha de Expediente (DD/MM/YYYY)"),
            ("LINK_TUC", "R", False, "Enlace a documento TUC en Drive"),
            ("LINK_NOTIFICACION", "S", False, "Enlace a Notificación en Drive"),
            ("DETALLES", "T", False, "Detalles adicionales"),
        ]

        for col_idx, (header, _, _req, _desc) in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
            cell.alignment = center
            col_letter = chr(64 + col_idx) if col_idx <= 26 else f"A{chr(64 + col_idx - 26)}"
            ws.column_dimensions[col_letter].width = 18

        for col_idx, (_, _, _, desc) in enumerate(headers, 1):
            cell = ws.cell(row=2, column=col_idx, value=desc)
            cell.font = Font(size=8, italic=True, color="555555")
            cell.alignment = Alignment(wrap_text=True)
            ws.row_dimensions[2].height = 35

        ejemplos = [
            "20123456789", "R-0128-2024", "0133-2024-S", "100%",
            "A2B-123", "01,02,03", "T-010145",
            "HABILITADO", "Sustitución aprobada | Vehículo nuevo",
            "23/04/2026", "EMPRESA EJEMPLO S.R.L.", "10/12/2025",
            "001", "SÍ", "ACTIVA", "EXP-2026-01290", "15/01/2026",
            "https://drive.google.com/...", "https://drive.google.com/...", "Trámite completado"
        ]
        for col_idx, ej in enumerate(ejemplos, 1):
            cell = ws.cell(row=3, column=col_idx, value=ej)
            cell.font = Font(size=9, color="0D6EFD")
            cell.alignment = Alignment(wrap_text=True)

        ws.row_dimensions[1].height = 30
        ws.freeze_panes = "A4"

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer
