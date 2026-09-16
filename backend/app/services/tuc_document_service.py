import os
import re
import io
import logging
from datetime import datetime, date
from typing import Dict, Any, Optional, List
from bson import ObjectId
import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.dependencies.db import get_database

logger = logging.getLogger(__name__)

TEMPLATE_PATH = os.path.join(os.path.dirname(os.path.dirname(__file__)), "templates", "plantilla_tuc.docx")

async def _get_db():
    database = await get_database()
    if database is None:
        raise RuntimeError("No se pudo conectar a la base de datos MongoDB")
    return database

def format_fecha(val: Any) -> str:
    if not val:
        return "-"
    if isinstance(val, (datetime, date)):
        return val.strftime("%d/%m/%Y")
    if isinstance(val, str):
        val_s = val.strip()
        try:
            dt = datetime.fromisoformat(val_s.replace("Z", "+00:00"))
            return dt.strftime("%d/%m/%Y")
        except Exception:
            m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", val_s)
            if m:
                return f"{m.group(3)}/{m.group(2)}/{m.group(1)}"
            return val_s
    return str(val)

def clean_num_resolucion(res_str: Optional[str]) -> str:
    if not res_str:
        return "-"
    s = res_str.strip().upper()
    if s.startswith("R-"):
        s = s[2:].strip()
    return s

class TucDocumentService:

    @staticmethod
    def _build_routes_table(rutas_info: List[Dict[str, Any]]):
        """
        Construye el elemento XML de Table 1 (Tabla de Rutas del reverso de la TUC)
        con la disposición, anchos y fuentes exactos de la plantilla oficial de DRTC Puno.
        """
        tbl_xml = [
            '<w:tbl ' + nsdecls('w') + '>',
            '  <w:tblPr>',
            '    <w:tblStyle w:val="Table2"/>',
            '    <w:tblW w:w="9210.0" w:type="dxa"/>',
            '    <w:jc w:val="left"/>',
            '    <w:tblBorders>',
            '      <w:top w:color="ffffff" w:space="0" w:sz="8" w:val="single"/>',
            '      <w:left w:color="ffffff" w:space="0" w:sz="8" w:val="single"/>',
            '      <w:bottom w:color="ffffff" w:space="0" w:sz="8" w:val="single"/>',
            '      <w:right w:color="ffffff" w:space="0" w:sz="8" w:val="single"/>',
            '      <w:insideH w:color="ffffff" w:space="0" w:sz="8" w:val="single"/>',
            '      <w:insideV w:color="ffffff" w:space="0" w:sz="8" w:val="single"/>',
            '    </w:tblBorders>',
            '  </w:tblPr>',
            '  <w:tblGrid>',
            '    <w:gridCol w:w="1980"/>',
            '    <w:gridCol w:w="4216"/>',
            '    <w:gridCol w:w="1410"/>',
            '    <w:gridCol w:w="1604"/>',
            '  </w:tblGrid>',
            '  <w:tr>',
            '    <w:trPr><w:cantSplit w:val="0"/><w:tblHeader w:val="0"/></w:trPr>',
            '    <w:tc><w:tcPr><w:tcW w:w="1980" w:type="dxa"/><w:vAlign w:val="top"/></w:tcPr><w:p><w:pPr><w:spacing w:after="0" w:before="0" w:line="192" w:lineRule="auto"/><w:rPr><w:rFonts w:ascii="Roboto" w:hAnsi="Roboto"/><w:sz w:val="14"/></w:rPr></w:pPr></w:p></w:tc>',
            '    <w:tc><w:tcPr><w:tcW w:w="4216" w:type="dxa"/><w:vAlign w:val="top"/></w:tcPr><w:p><w:pPr><w:spacing w:after="0" w:before="0" w:line="192" w:lineRule="auto"/><w:rPr><w:rFonts w:ascii="Roboto" w:hAnsi="Roboto"/><w:sz w:val="14"/></w:rPr></w:pPr></w:p></w:tc>',
            '    <w:tc><w:tcPr><w:tcW w:w="1410" w:type="dxa"/><w:vAlign w:val="top"/></w:tcPr><w:p><w:pPr><w:spacing w:after="0" w:before="0" w:line="192" w:lineRule="auto"/><w:rPr><w:rFonts w:ascii="Roboto" w:hAnsi="Roboto"/><w:sz w:val="14"/></w:rPr></w:pPr></w:p></w:tc>',
            '    <w:tc><w:tcPr><w:tcW w:w="1604" w:type="dxa"/><w:vAlign w:val="top"/></w:tcPr><w:p><w:pPr><w:spacing w:after="0" w:before="0" w:line="192" w:lineRule="auto"/><w:rPr><w:rFonts w:ascii="Roboto" w:hAnsi="Roboto"/><w:sz w:val="14"/></w:rPr></w:pPr></w:p></w:tc>',
            '  </w:tr>'
        ]

        lista = list(rutas_info) if rutas_info else [{"codigo": "", "tramo": "SIN RUTAS ASIGNADAS", "frecuencia": ""}]

        for r in lista:
            cod_str = f"Ruta {r.get('codigo')}: " if r.get('codigo') else ""
            tramo_str = str(r.get('tramo') or f"{r.get('origen', '')} - {r.get('destino', '')}").strip()
            frec_str = str(r.get('frecuencia', '')).strip()

            # Sanitizar para XML
            cod_str = cod_str.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            tramo_str = tramo_str.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            frec_str = frec_str.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

            row_xml = f'''
              <w:tr>
                <w:trPr><w:cantSplit w:val="0"/><w:tblHeader w:val="0"/></w:trPr>
                <w:tc>
                  <w:tcPr><w:tcW w:w="1980" w:type="dxa"/><w:vAlign w:val="top"/></w:tcPr>
                  <w:p><w:pPr><w:spacing w:after="0" w:before="0" w:line="192" w:lineRule="auto"/><w:rPr><w:rFonts w:ascii="Roboto" w:hAnsi="Roboto"/><w:sz w:val="14"/></w:rPr></w:pPr></w:p>
                </w:tc>
                <w:tc>
                  <w:tcPr><w:tcW w:w="4216" w:type="dxa"/><w:vAlign w:val="top"/></w:tcPr>
                  <w:p>
                    <w:pPr>
                      <w:spacing w:after="0" w:before="0" w:line="192" w:lineRule="auto"/>
                      <w:rPr><w:rFonts w:ascii="Roboto" w:hAnsi="Roboto"/><w:sz w:val="14"/></w:rPr>
                    </w:pPr>
                    <w:r>
                      <w:rPr><w:rFonts w:ascii="Roboto" w:hAnsi="Roboto"/><w:b w:val="1"/><w:sz w:val="14"/></w:rPr>
                      <w:t xml:space="preserve">{cod_str}</w:t>
                    </w:r>
                    <w:r>
                      <w:rPr><w:rFonts w:ascii="Roboto" w:hAnsi="Roboto"/><w:b w:val="0"/><w:sz w:val="14"/></w:rPr>
                      <w:t xml:space="preserve">{tramo_str}</w:t>
                    </w:r>
                  </w:p>
                </w:tc>
                <w:tc>
                  <w:tcPr><w:tcW w:w="1410" w:type="dxa"/><w:vAlign w:val="top"/></w:tcPr>
                  <w:p>
                    <w:pPr>
                      <w:spacing w:after="0" w:before="0" w:line="192" w:lineRule="auto"/>
                      <w:rPr><w:rFonts w:ascii="Roboto" w:hAnsi="Roboto"/><w:sz w:val="14"/></w:rPr>
                    </w:pPr>
                    <w:r>
                      <w:rPr><w:rFonts w:ascii="Roboto" w:hAnsi="Roboto"/><w:b w:val="0"/><w:sz w:val="14"/></w:rPr>
                      <w:t xml:space="preserve">{frec_str}</w:t>
                    </w:r>
                  </w:p>
                </w:tc>
                <w:tc>
                  <w:tcPr><w:tcW w:w="1604" w:type="dxa"/><w:vAlign w:val="top"/></w:tcPr>
                  <w:p><w:pPr><w:spacing w:after="0" w:before="0" w:line="192" w:lineRule="auto"/><w:rPr><w:rFonts w:ascii="Roboto" w:hAnsi="Roboto"/><w:sz w:val="14"/></w:rPr></w:pPr></w:p>
                </w:tc>
              </w:tr>
            '''
            tbl_xml.append(row_xml)

        tbl_xml.append('</w:tbl>')
        return parse_xml('\n'.join(tbl_xml))

    @staticmethod
    async def get_tuc_data(placa_o_id: str) -> Dict[str, Any]:
        """
        Recopila y cruza toda la información técnica y administrativa
        de un vehículo para rellenar la plantilla oficial de TUC.
        """
        db = await _get_db()
        term = placa_o_id.strip()

        # 1. Buscar vehículo en flota_empresa
        query: Dict[str, Any] = {}
        if ObjectId.is_valid(term):
            query = {"_id": ObjectId(term)}
        else:
            query = {
                "$or": [
                    {"id": term},
                    {"placa": {"$regex": f"^{re.escape(term)}$", "$options": "i"}},
                    {"numero_tuc": {"$regex": f"^{re.escape(term)}$", "$options": "i"}}
                ]
            }

        vehiculo = await db.flota_empresa.find_one(query)
        if not vehiculo:
            placa_limpia = term.replace("-", "").strip()
            vehiculo = await db.flota_empresa.find_one({
                "placa": {"$regex": f"^{re.escape(placa_limpia)}$", "$options": "i"}
            })

        if not vehiculo:
            raise ValueError(f"No se encontró el vehículo con identificador o placa '{placa_o_id}' en la flota.")

        placa = (vehiculo.get("placa") or "").strip().upper()
        ruc = (vehiculo.get("ruc") or "").strip()
        razon_social = (vehiculo.get("razon_social") or "").strip()
        nro_primigenia_raw = (vehiculo.get("nro_resolucion_primigenia") or "").strip()
        nro_hija_raw = (vehiculo.get("nro_resolucion_hija") or "").strip()
        tipo_hija = (vehiculo.get("tipo_resolucion_hija") or "").strip()
        fecha_hija = vehiculo.get("fecha_resolucion_hija")
        rutas_codigos = vehiculo.get("rutas") or []

        # 2. Obtener datos técnicos de vehiculos_data
        v_data = await db.vehiculos_data.find_one({
            "$or": [
                {"placa_actual": {"$regex": f"^{re.escape(placa)}$", "$options": "i"}},
                {"placa": {"$regex": f"^{re.escape(placa)}$", "$options": "i"}}
            ]
        })
        if not v_data and "-" in placa:
            placa_alt = placa.replace("-", "")
            v_data = await db.vehiculos_data.find_one({
                "$or": [
                    {"placa_actual": {"$regex": f"^{re.escape(placa_alt)}$", "$options": "i"}},
                    {"placa": {"$regex": f"^{re.escape(placa_alt)}$", "$options": "i"}}
                ]
            })

        color = (v_data.get("color") or "-").strip().upper() if v_data else "-"
        marca = (v_data.get("marca") or "-").strip().upper() if v_data else "-"
        vin = (v_data.get("vin") or v_data.get("numero_serie") or "-").strip().upper() if v_data else "-"
        anio = str(v_data.get("anio_fabricacion") or v_data.get("anio_modelo") or "-") if v_data else "-"
        asientos = str(v_data.get("numero_asientos") or "-") if v_data else "-"
        alto = str(v_data.get("altura") or "-") if v_data else "-"
        ancho = str(v_data.get("ancho") or "-") if v_data else "-"
        largo = str(v_data.get("longitud") or "-") if v_data else "-"
        peso_neto = str(v_data.get("peso_neto") or v_data.get("peso_seco") or "-") if v_data else "-"
        peso_bruto = str(v_data.get("peso_bruto") or "-") if v_data else "-"
        carga_util = str(v_data.get("carga_util") or "-") if v_data else "-"
        categoria = (v_data.get("categoria") or "-").strip().upper() if v_data else "-"
        ejes = str(v_data.get("numero_ejes") or "-") if v_data else "-"

        # 3. Obtener datos de la empresa (Partida Registral)
        empresa = None
        if ruc:
            empresa = await db.empresas.find_one({"ruc": ruc})
        if not empresa and razon_social:
            empresa = await db.empresas.find_one({"razonSocial.principal": {"$regex": f"^{re.escape(razon_social)}$", "$options": "i"}})

        partida = "-"
        if empresa:
            partida = (
                empresa.get("partidaRegistral") or 
                empresa.get("partida_registral") or 
                empresa.get("partida") or 
                (empresa.get("datosSunat") or {}).get("ddp_numreg") or
                "-"
            )
            if not razon_social:
                rs = empresa.get("razonSocial")
                if isinstance(rs, dict):
                    razon_social = rs.get("principal") or rs.get("sunat") or ""
                elif isinstance(rs, str):
                    razon_social = rs

        if (not partida or partida == "-") and v_data:
            partida = v_data.get("partida_registral") or "-"

        # 4. Obtener datos de la resolución primigenia
        res_prim = None
        if nro_primigenia_raw:
            norm_res = clean_num_resolucion(nro_primigenia_raw)
            res_prim = await db.resoluciones_primigenias.find_one({
                "$or": [
                    {"nro_resolucion": {"$regex": f"^{re.escape(nro_primigenia_raw)}$", "$options": "i"}},
                    {"nro_resolucion": {"$regex": f"{re.escape(norm_res)}$", "$options": "i"}}
                ]
            })

        fecha_del = "-"
        fecha_al = "-"
        fecha_res_p = "-"
        nro_res_p = clean_num_resolucion(nro_primigenia_raw)

        if res_prim:
            fecha_del = format_fecha(res_prim.get("fecha_inicio_vigencia") or res_prim.get("fecha_resolucion"))
            fecha_al = format_fecha(res_prim.get("fecha_fin_vigencia"))
            fecha_res_p = format_fecha(res_prim.get("fecha_resolucion") or res_prim.get("fecha_emision"))
            if not nro_res_p or nro_res_p == "-":
                nro_res_p = clean_num_resolucion(res_prim.get("nro_resolucion"))
        else:
            fecha_al = format_fecha(vehiculo.get("fecha_vigencia_hasta"))

        # 5. Obtener detalle de las rutas asociadas a la empresa y al vehículo
        rutas_objs = []
        cods_limpios = [str(c).strip() for c in rutas_codigos if c and str(c).strip()] if isinstance(rutas_codigos, list) else []

        # Búsqueda 1: Por RUC de empresa y códigos específicos del vehículo
        if cods_limpios and ruc:
            rutas_cursor = db.rutas.find({
                "empresa.ruc": ruc,
                "codigoRuta": {"$in": cods_limpios}
            })
            rutas_objs = await rutas_cursor.to_list(100)

        # Búsqueda 2: Si no hubo coincidencia con RUC, buscar códigos en general
        if not rutas_objs and cods_limpios:
            rutas_cursor = db.rutas.find({"codigoRuta": {"$in": cods_limpios}})
            rutas_objs = await rutas_cursor.to_list(100)

        # Búsqueda 3: Si el vehículo no tenía lista de rutas, tomar todas las de la empresa por su RUC
        if not rutas_objs and ruc:
            rutas_cursor = db.rutas.find({"empresa.ruc": ruc})
            rutas_objs = await rutas_cursor.to_list(100)

        # Búsqueda 4: Por número de resolución primigenia en resolucion.nroResolucion
        if not rutas_objs and nro_primigenia_raw:
            clean_res = clean_num_resolucion(nro_primigenia_raw)
            rutas_cursor = db.rutas.find({"resolucion.nroResolucion": {"$regex": re.escape(clean_res), "$options": "i"}})
            rutas_objs = await rutas_cursor.to_list(100)

        rutas_map = {str(r.get("codigoRuta", "")).strip(): r for r in rutas_objs}
        rutas_lineas = []
        rutas_para_frontend = []

        # Determinar lista de códigos ordenados a procesar
        if cods_limpios:
            cods_a_procesar = sorted(cods_limpios)
        elif rutas_objs:
            cods_a_procesar = sorted([str(r.get("codigoRuta", "")).strip() for r in rutas_objs if r.get("codigoRuta")])
        else:
            cods_a_procesar = []

        for cod in cods_a_procesar:
            if not cod:
                continue
            r_obj = rutas_map.get(cod)
            if r_obj:
                origen = r_obj.get("origen")
                origen_nom = origen.get("nombre") if isinstance(origen, dict) else str(origen or "")

                destino = r_obj.get("destino")
                destino_nom = destino.get("nombre") if isinstance(destino, dict) else str(destino or "")

                itinerario = r_obj.get("itinerario")
                itin_items = []
                if isinstance(itinerario, list):
                    for it in itinerario:
                        if isinstance(it, dict):
                            nom = it.get("nombre") or it.get("descripcion") or ""
                            if nom:
                                itin_items.append(str(nom))
                        elif isinstance(it, str) and it.strip():
                            itin_items.append(it.strip())
                elif isinstance(itinerario, str) and itinerario.strip():
                    itin_items.append(itinerario.strip())
                itin_str = " - ".join(itin_items)

                frecuencia = r_obj.get("frecuencia")
                frec = ""
                if isinstance(frecuencia, dict):
                    frec = frecuencia.get("descripcion") or frecuencia.get("tipo") or ""
                elif isinstance(frecuencia, str):
                    frec = frecuencia

                partes_tramo = [p for p in [origen_nom, itin_str, destino_nom] if p]
                tramo = " - ".join(partes_tramo)

                frec_display = f" ({frec})" if frec else ""
                rutas_lineas.append(f"Ruta {cod}: {tramo}{frec_display}")
                rutas_para_frontend.append({
                    "codigo": cod,
                    "origen": origen_nom,
                    "itinerario": itin_str,
                    "destino": destino_nom,
                    "tramo": tramo,
                    "frecuencia": frec
                })
            else:
                rutas_lineas.append(f"Ruta {cod}")
                rutas_para_frontend.append({
                    "codigo": cod,
                    "origen": "",
                    "itinerario": "",
                    "destino": "",
                    "tramo": f"Ruta {cod}",
                    "frecuencia": ""
                })

        tabla_rutas_text = "\n".join(rutas_lineas) if rutas_lineas else "SIN RUTAS ASIGNADAS"

        # 6. Acto resolutivo final (Hija o Primigenia)
        num_resolucion_final = clean_num_resolucion(nro_hija_raw or nro_primigenia_raw)
        fecha_res_final = format_fecha(fecha_hija or fecha_res_p)
        tipo_res_final = tipo_hija.upper() if tipo_hija else "AUTORIZACION"
        if not tipo_hija and res_prim and res_prim.get("tipo_autorizacion"):
            tipo_res_final = res_prim.get("tipo_autorizacion").upper()

        # Mapeo completo de las 25 etiquetas oficiales
        placeholders = {
            "{{FECHA_DEL}}": fecha_del,
            "{{FECHA_AL}}": fecha_al,
            "{{RES}}": nro_res_p,
            "{{FECHA_RES_P}}": fecha_res_p,
            "{{EMPRESA}}": razon_social,
            "{{RUC}}": ruc,
            "{{PARTIDA}}": partida,
            "{{PLACA}}": placa,
            "{{COLOR}}": color,
            "{{MARCA}}": marca,
            "{{VIN}}": vin,
            "{{ANIO}}": anio,
            "{{ASIENTOS}}": asientos,
            "{{ALTO}}": alto,
            "{{PESO_NETO}}": peso_neto,
            "{{CATEGORIA}}": categoria,
            "{{EJES}}": ejes,
            "{{ANCHO}}": ancho,
            "{{CARGA_UTIL}}": carga_util,
            "{{LARGO}}": largo,
            "{{PESO_BRUTO}}": peso_bruto,
            "{{TABLA_RUTAS}}": tabla_rutas_text,
            "{{NUM_RESOLUCION}}": num_resolucion_final,
            "{{FECHA_RES}}": fecha_res_final,
            "{{TIPO_RES}}": tipo_res_final
        }

        datos_estructurados = {
            "placa": placa,
            "numero_tuc": vehiculo.get("numero_tuc") or vehiculo.get("tuc") or "",
            "ruc": ruc,
            "empresa": razon_social,
            "partida": partida,
            "fecha_del": fecha_del,
            "fecha_al": fecha_al,
            "nro_resolucion_primigenia": nro_res_p,
            "fecha_resolucion_primigenia": fecha_res_p,
            "color": color,
            "marca": marca,
            "vin": vin,
            "anio": anio,
            "asientos": asientos,
            "alto": alto,
            "ancho": ancho,
            "largo": largo,
            "peso_neto": peso_neto,
            "peso_bruto": peso_bruto,
            "carga_util": carga_util,
            "categoria": categoria,
            "ejes": ejes,
            "rutas_detalle": rutas_para_frontend,
            "tabla_rutas_text": tabla_rutas_text,
            "num_resolucion_acto": num_resolucion_final,
            "fecha_resolucion_acto": fecha_res_final,
            "tipo_resolucion_acto": tipo_res_final
        }

        return {
            "placeholders": placeholders,
            "datos_estructurados": datos_estructurados,
            "placa": placa,
            "numero_tuc": vehiculo.get("numero_tuc") or vehiculo.get("tuc") or ""
        }

    @staticmethod
    def _reemplazar_en_parrafo(p, old_text: str, new_text: str):
        if old_text not in p.text:
            return
        full_text = p.text.replace(old_text, str(new_text))
        if p.runs:
            p.runs[0].text = full_text
            for r in p.runs[1:]:
                r.text = ""
        else:
            p.text = full_text

    @staticmethod
    async def generar_docx_tuc(placa_o_id: str) -> Dict[str, Any]:
        """
        Abre la plantilla oficial .docx y sustituye los campos con la data real,
        insertando la tabla de rutas estructurada con 4 columnas y bordes transparentes.
        """
        if not os.path.exists(TEMPLATE_PATH):
            raise FileNotFoundError(f"No se encontró la plantilla oficial de TUC en {TEMPLATE_PATH}")

        tuc_info = await TucDocumentService.get_tuc_data(placa_o_id)
        placeholders = tuc_info["placeholders"]
        placa = tuc_info["placa"]
        rutas_detalle = tuc_info["datos_estructurados"].get("rutas_detalle", [])

        doc = docx.Document(TEMPLATE_PATH)

        # 1. Insertar la tabla de rutas en lugar del marcador {{TABLA_RUTAS}}
        for p in list(doc.paragraphs):
            if "{{TABLA_RUTAS}}" in p.text:
                tbl_el = TucDocumentService._build_routes_table(rutas_detalle)
                p._p.addprevious(tbl_el)
                p._p.getparent().remove(p._p)
                break

        # 2. Reemplazar los demás marcadores en párrafos restantes
        for p in doc.paragraphs:
            for k, v in placeholders.items():
                if k == "{{TABLA_RUTAS}}":
                    continue
                TucDocumentService._reemplazar_en_parrafo(p, k, v)

        # 3. Reemplazar en tablas (Ficha técnica anverso)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for k, v in placeholders.items():
                            if k == "{{TABLA_RUTAS}}":
                                continue
                            TucDocumentService._reemplazar_en_parrafo(p, k, v)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        filename = f"TUC_{placa.replace('-', '_')}.docx"
        return {
            "buffer": buffer,
            "filename": filename,
            "placa": placa,
            "datos": tuc_info["datos_estructurados"]
        }

    @staticmethod
    async def get_datos_impresion(placa_o_id: str) -> Dict[str, Any]:
        """
        Retorna la data completa y estructurada lista para la vista previa de impresión
        HTML/CSS directamente en el navegador del operador.
        """
        tuc_info = await TucDocumentService.get_tuc_data(placa_o_id)
        return {
            "placa": tuc_info["placa"],
            "numero_tuc": tuc_info["numero_tuc"],
            "datos": tuc_info["datos_estructurados"],
            "placeholders": tuc_info["placeholders"]
        }
